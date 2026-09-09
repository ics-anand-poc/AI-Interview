"""
Add/reset Employee Portal credentials for Swethasri R (1035294), CMG / PACO.
Does not wipe other portal users. Does not replace an in-progress/completed test.
"""
from __future__ import annotations

import sys
import uuid
from datetime import datetime, timezone
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))

from add_paco_cmg3_portal_users import (  # noqa: E402
    dump_existing,
    email_taken_by_other,
)
from add_paco_phase2_portal_users import (  # noqa: E402
    ACCOUNTS_FILE,
    CREDENTIALS_FILE,
    PROFILES_FILE,
    QB_FILE,
    QUESTIONS_PER_EMPLOYEE,
    SUBJECT_ID,
    SUBJECT_TITLE,
    TIME_LIMIT_SECONDS,
    append_local_tests,
    assign_display_questions,
    build_bank_index,
    clean,
    create_test_if_missing,
    generate_password,
    hash_password,
    load_json,
    normalize_question_key,
    parse_qb_new_xlsx,
    save_json,
    upsert_supabase_employee,
)
from reassign_test_questions_supabase import TOPIC_ID, SupabaseClient, load_env  # noqa: E402

EMP = {
    "employee_id": "1035294",
    "full_name": "Swethasri R",
    "email": "swethasri.ravikumar.ext@nokia.com",
    "infinite_email": "Swethasri.Ravikumar@infinite.com",
    "domain": "PACO",
    "product": "CMG",
    "role": "employee",
    "ddh": "",
}

SHARE_FILE = Path.home() / "OneDrive" / "Desktop" / "Swethasri_R_CMG_Credentials.xlsx"


def upsert_credentials_row(emp: dict) -> None:
    if CREDENTIALS_FILE.exists():
        wb = openpyxl.load_workbook(CREDENTIALS_FILE)
        ws = wb.active
        header = [clean(c.value) for c in ws[1]]
        col = {h: i + 1 for i, h in enumerate(header)}
        emp_col = col.get("Emp ID", 1)
        target = None
        for r in range(2, ws.max_row + 1):
            if clean(ws.cell(r, emp_col).value).upper() == emp["employee_id"].upper():
                target = r
                break
        if target is None:
            target = ws.max_row + 1
    else:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "User Credentials"
        header = [
            "Emp ID",
            "Employee Name",
            "Initial Password",
            "Nokia Email ID",
            "Infinite Email ID",
            "Role",
            "Domain",
            "Product",
            "Customer",
            "DDH Manager",
        ]
        ws.append(header)
        col = {h: i + 1 for i, h in enumerate(header)}
        emp_col = 1
        target = 2

    values = {
        "Emp ID": emp["employee_id"],
        "Employee Name": emp["full_name"],
        "Initial Password": emp["password"],
        "Nokia Email ID": emp["email"] if "@nokia.com" in emp["email"].lower() else EMP["email"],
        "Infinite Email ID": emp.get("infinite_email") or "",
        "Role": emp["role"],
        "Domain": emp["domain"],
        "Product": emp["product"],
        "Customer": "",
        "DDH Manager": emp["ddh"],
    }
    ws.cell(target, emp_col).value = str(emp["employee_id"])
    for key, value in values.items():
        if key in col:
            ws.cell(target, col[key]).value = value
    CREDENTIALS_FILE.parent.mkdir(parents=True, exist_ok=True)
    wb.save(CREDENTIALS_FILE)

    share = openpyxl.Workbook()
    sws = share.active
    sws.title = "Credentials"
    sws.append(
        [
            "Emp ID",
            "Employee Name",
            "Initial Password",
            "Infinite Email",
            "Nokia Email",
            "Product",
            "Domain",
            "Login URL",
        ]
    )
    sws.append(
        [
            emp["employee_id"],
            emp["full_name"],
            emp["password"],
            emp.get("infinite_email") or "",
            EMP["email"],
            emp["product"],
            emp["domain"],
            "https://ai-interview-ics-poc.vercel.app/employee",
        ]
    )
    share.save(SHARE_FILE)


def main() -> int:
    if not QB_FILE.exists():
        print(f"ERROR: missing {QB_FILE}")
        return 1

    url, key = load_env()
    client = SupabaseClient(url, key)
    emp = dict(EMP)

    print("=== Existing Supabase row ===")
    dump_existing(client, emp["employee_id"], emp["email"])
    dump_existing(client, emp["employee_id"], emp["infinite_email"])

    pools, mcq_records = parse_qb_new_xlsx(QB_FILE)
    bank = build_bank_index(mcq_records)
    now = datetime.now(timezone.utc).isoformat()

    product = emp["product"]
    pool = pools.get(product, [])
    assigned, remark = assign_display_questions(product, pool, employee_id=emp["employee_id"])
    assigned = [q for q in assigned if q]
    if len(assigned) < QUESTIONS_PER_EMPLOYEE:
        print(f"ERROR: only got {len(assigned)} {product} questions. {remark}")
        return 1

    matched = []
    for q_text in assigned:
        item = bank.get(normalize_question_key(q_text)) or bank.get(
            normalize_question_key(q_text.split("] ", 1)[-1])
        )
        if not item:
            print(f"ERROR: unmatched question: {q_text[:120]}")
            return 1
        matched.append({**item, "question_text": q_text})

    emp["password"] = generate_password(emp["full_name"], emp["employee_id"])
    emp["password_hash"], emp["password_salt"] = hash_password(emp["password"])

    login_email = emp["infinite_email"]
    if email_taken_by_other(client, login_email, emp["employee_id"]):
        print(f"  Infinite email already on another Emp ID; using Nokia email")
        login_email = emp["email"]
    elif email_taken_by_other(client, emp["email"], emp["employee_id"]):
        print(f"  Nokia email already on another Emp ID; using Infinite email")
        login_email = emp["infinite_email"]
    emp["account_email"] = login_email

    accounts = load_json(ACCOUNTS_FILE, {"employees": []})
    account_map = {
        clean(acc.get("employee_id")).upper(): acc
        for acc in accounts.get("employees", [])
        if acc.get("employee_id")
    }
    existing_acc = account_map.get(emp["employee_id"].upper(), {})
    account = {
        **existing_acc,
        "employee_id": emp["employee_id"],
        "full_name": emp["full_name"],
        "email": login_email,
        "department": emp["domain"],
        "role": emp["role"],
        "product": product,
        "assessment_only": True,
        "product_qb_eligible": True,
        "is_first_login": True,
        "password_hash": emp["password_hash"],
        "password_salt": emp["password_salt"],
        "xp_points": existing_acc.get("xp_points", 0),
        "streak_days": existing_acc.get("streak_days", 0),
        "skill_level": existing_acc.get("skill_level", "beginner"),
        "ai_readiness_score": existing_acc.get("ai_readiness_score", 0),
    }
    account_map[emp["employee_id"].upper()] = account

    profiles = load_json(PROFILES_FILE, [])
    profile_map = {clean(row.get("employee_id")).upper(): row for row in profiles}
    profile_map[emp["employee_id"].upper()] = {
        "employee_id": emp["employee_id"],
        "full_name": emp["full_name"],
        "role": emp["role"],
        "domain": emp["domain"],
        "product": product,
        "email": login_email,
        "ddh": emp["ddh"],
        "emp_status": "Confirmed",
        "remarks": "PACO CMG portal credentials",
        "assigned_questions": assigned,
        "assigned_question_count": len(assigned),
    }

    test_id = str(uuid.uuid4())
    new_tests = [
        {
            "id": test_id,
            "employee_id": emp["employee_id"],
            "employee_code": emp["employee_id"],
            "topic_id": TOPIC_ID,
            "subject_id": SUBJECT_ID,
            "difficulty": "medium",
            "total_questions": len(matched),
            "time_limit_seconds": TIME_LIMIT_SECONDS,
            "status": "pending",
            "current_question_index": 0,
            "started_at": None,
            "completed_at": None,
            "in_progress": None,
            "created_at": now,
            "topic_title": product,
            "subject_title": SUBJECT_TITLE,
        }
    ]
    new_questions = [
        {
            "id": str(uuid.uuid4()),
            "test_id": test_id,
            "question_index": idx,
            "question_text": q["question_text"],
            "options": q["options"],
            "correct_option_index": q["correct_option_index"],
            "explanation": q["explanation"],
            "difficulty": q["difficulty"],
            "topic_id": TOPIC_ID,
            "topic_title": q.get("category") or q["topic_title"],
            "created_at": now,
        }
        for idx, q in enumerate(matched)
    ]
    new_manifest = [
        {
            "employee_id": emp["employee_id"],
            "full_name": emp["full_name"],
            "product": product,
            "test_id": test_id,
            "question_count": len(matched),
            "missing_questions": [],
        }
    ]

    accounts["employees"] = list(account_map.values())
    save_json(ACCOUNTS_FILE, accounts)
    save_json(PROFILES_FILE, list(profile_map.values()))
    upsert_credentials_row(emp)
    append_local_tests(new_tests, new_questions, new_manifest)

    print("\n=== Syncing to Supabase Employee Portal ===")
    employee_uuid = upsert_supabase_employee(client, emp, account)
    create_test_if_missing(client, employee_uuid, emp, matched, test_id)
    print(f"  synced {emp['employee_id']} {emp['full_name']}")

    print("\n=== Employee Portal login ===")
    print("URL: https://ai-interview-ics-poc.vercel.app/employee")
    print(f"Emp ID:   {emp['employee_id']}")
    print(f"Name:     {emp['full_name']}")
    print(f"Email:    {login_email}")
    print(f"Nokia:    {EMP['email']}")
    print(f"Product:  {product}")
    print(f"Password: {emp['password']}")
    print(f"Share:    {SHARE_FILE}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
