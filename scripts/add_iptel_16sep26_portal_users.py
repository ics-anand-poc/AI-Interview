"""
Create Employee Portal credentials for IPTEL people in
'Employee crediantials to be created - IPTEL.xlsx'.

Assigns Q-forms from QB-new.xlsx for NN, NTAS, SBC, CFX, NEF.
MSS/MGW has no QB bank yet — login is still created.
Writes a new share file with passwords. Does not wipe other portal users.
"""
from __future__ import annotations

import re
import sys
import uuid
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))

from add_paco_cmg3_portal_users import email_taken_by_other, replace_or_create_test  # noqa: E402
from add_paco_16sep26_portal_users import create_test_if_missing, names_match  # noqa: E402
from add_paco_phase2_portal_users import (  # noqa: E402
    ACCOUNTS_FILE,
    PROFILES_FILE,
    QUESTIONS_PER_EMPLOYEE,
    SUBJECT_ID,
    SUBJECT_TITLE,
    TIME_LIMIT_SECONDS,
    append_credentials_rows,
    append_local_tests,
    assign_display_questions,
    build_bank_index,
    clean,
    generate_password,
    hash_password,
    load_existing_passwords,
    load_json,
    normalize_question_key,
    parse_qb_new_xlsx,
    save_json,
    update_mapping_row,
    upsert_supabase_employee,
)
import add_paco_phase2_portal_users as phase2  # noqa: E402
from qb_new_parser import resolve_qb_product_key  # noqa: E402
from reassign_test_questions_supabase import TOPIC_ID, SupabaseClient, load_env  # noqa: E402

DOCS = next(
    (
        p
        for p in (ROOT / "NON-Needed docs", ROOT / "docs" / "NON-Needed docs")
        if (p / "QB-new.xlsx").exists()
    ),
    ROOT / "docs" / "NON-Needed docs",
)
QB_FILE = DOCS / "QB-new.xlsx"
SOURCE_FILE = ROOT / "Employee crediantials to be created - IPTEL.xlsx"
SHARE_FILE = DOCS / "IPTEL_16Sep26_User_Credentials.xlsx"
SHARE_FILE_ROOT = ROOT / "IPTEL_16Sep26_User_Credentials.xlsx"
REMARKS = "IPTEL credentials 16 Sep 2026"
phase2.QB_FILE = QB_FILE
phase2.CREDENTIALS_FILE = DOCS / "Employee_User_Credentials.xlsx"
phase2.MAPPING_FILE = DOCS / "Resource_Question_Mapping.xlsx"
phase2.DOCS = DOCS


def load_source_employees() -> list[dict]:
    if not SOURCE_FILE.exists():
        raise FileNotFoundError(SOURCE_FILE)
    rows = list(load_workbook(SOURCE_FILE, data_only=True).active.iter_rows(values_only=True))
    header = [clean(c) for c in rows[0]]
    col = {name.lower(): idx for idx, name in enumerate(header)}

    def get(row, *keys):
        for key in keys:
            idx = col.get(key.lower())
            if idx is None or idx >= len(row):
                continue
            value = clean(row[idx]).replace("\xa0", " ")
            if value and value.upper() not in {"#N/A", "N/A", "NA", "NONE", "NULL"}:
                return value
        return ""

    people = []
    seen = set()
    for row in rows[1:]:
        emp_id = get(row, "Emp ID")
        if not emp_id or not re.search(r"\d", emp_id) or emp_id.upper() in seen:
            continue
        seen.add(emp_id.upper())
        raw_product = get(row, "Primary Product", "Product")
        people.append(
            {
                "employee_id": emp_id,
                "full_name": get(row, "Emp Name"),
                "email": get(row, "Nokia Email ID", "Email"),
                "role": get(row, "Role") or "employee",
                "domain": get(row, "Domain") or "IPTEL",
                "product": resolve_qb_product_key(raw_product) or raw_product,
                "raw_product": raw_product,
                "project": get(row, "Current Project"),
                "ddh": get(row, "DDH"),
            }
        )
    return people


def write_share_file(employees: list[dict]) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "IPTEL 16 Sep 26"
    header = [
        "Emp ID",
        "Employee Name",
        "Initial Password",
        "Nokia Email ID",
        "Role",
        "Domain",
        "Product",
        "Current Project",
        "DDH",
        "Login ID",
        "Portal status",
    ]
    ws.append(header)
    for emp in employees:
        ws.append(
            [
                emp["employee_id"],
                emp["full_name"],
                emp["password"],
                emp["email"],
                emp.get("role") or "employee",
                emp.get("domain") or "IPTEL",
                emp.get("raw_product") or emp["product"],
                emp.get("project") or "",
                emp.get("ddh") or "",
                emp["employee_id"],
                emp["portal_status"],
            ]
        )
        row = ws.max_row
        ws.cell(row, 1).value = str(emp["employee_id"])
        ws.cell(row, 10).value = str(emp["employee_id"])
        ws.cell(row, 3).font = Font(name="Consolas", size=10, bold=True, color="4338CA")

    header_fill = PatternFill("solid", fgColor="4F46E5")
    for cell in ws[1]:
        cell.font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(vertical="center", horizontal="center")
    ws.row_dimensions[1].height = 28
    widths = [14, 32, 24, 42, 28, 10, 12, 28, 18, 14, 36]
    for idx, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = width

    note = wb.create_sheet("How to login")
    note["A1"] = "Employee Portal login"
    note["A1"].font = Font(bold=True, size=14)
    note["A3"] = "URL"
    note["B3"] = "https://ai-interview-ics-poc.vercel.app/employee"
    note["A4"] = "Login ID"
    note["B4"] = "Emp ID from this sheet"
    note["A5"] = "Password"
    note["B5"] = "Initial Password column"
    note["A7"] = "Q-forms come from QB-new.xlsx (NN, NTAS, SBC, CFX, NEF). MSS/MGW has no question bank yet."
    note["A8"] = "Treat this file as confidential."
    note.column_dimensions["A"].width = 16
    note.column_dimensions["B"].width = 110
    SHARE_FILE.parent.mkdir(parents=True, exist_ok=True)
    wb.save(SHARE_FILE)
    wb.save(SHARE_FILE_ROOT)
    print(f"Wrote share file: {SHARE_FILE_ROOT} ({len(employees)} users)")


def main() -> int:
    if not QB_FILE.exists():
        print(f"ERROR: missing {QB_FILE}")
        return 1
    people = load_source_employees()
    if not people:
        print("ERROR: no people in IPTEL credentials workbook")
        return 1
    print(f"Loaded {len(people)} IPTEL people from {SOURCE_FILE.name}")

    stored_passwords = load_existing_passwords()
    accounts = load_json(ACCOUNTS_FILE, {"employees": []})
    account_map = {
        clean(acc.get("employee_id")).upper(): acc
        for acc in accounts.get("employees", [])
        if acc.get("employee_id")
    }
    profiles = load_json(PROFILES_FILE, [])
    profile_map = {clean(row.get("employee_id")).upper(): row for row in profiles}

    url, key = load_env()
    client = SupabaseClient(url, key)

    to_add = []
    share_rows = []
    for emp in people:
        existing_acc = account_map.get(emp["employee_id"].upper(), {})
        existing_name = clean(existing_acc.get("full_name"))
        eligible = bool(existing_acc.get("product_qb_eligible"))
        same_person = (not existing_name) or names_match(existing_name, emp["full_name"])
        if existing_name and not same_person:
            emp["action"] = "remap"
            emp["previous_name"] = existing_name
        elif eligible and same_person:
            emp["action"] = "skip"
            emp["previous_name"] = existing_name or emp["full_name"]
        else:
            emp["action"] = "add"
            emp["previous_name"] = existing_name or "(new)"

        if emp["action"] == "skip":
            emp["password"] = stored_passwords.get(emp["employee_id"].upper()) or "(already in portal — password unchanged)"
            emp["portal_status"] = "Already in portal"
            emp["skip"] = True
            share_rows.append(emp)
            print(f"SKIP {emp['employee_id']} {emp['full_name']} (already Q-form eligible)")
            continue

        emp["password"] = stored_passwords.get(emp["employee_id"].upper()) or generate_password(
            emp["full_name"], emp["employee_id"]
        )
        emp["password_hash"], emp["password_salt"] = hash_password(emp["password"])
        emp["portal_status"] = "Remapped Emp ID" if emp["action"] == "remap" else "Newly added"
        emp["skip"] = False
        to_add.append(emp)
        share_rows.append(emp)

    print(f"Already in portal: {sum(1 for e in share_rows if e['skip'])}")
    print(f"To add/remap: {len(to_add)}")

    pools, mcq_records = parse_qb_new_xlsx(QB_FILE)
    bank = build_bank_index(mcq_records)
    now = datetime.now(timezone.utc).isoformat()
    new_tests = []
    new_questions = []
    new_manifest = []

    for emp in to_add:
        product = emp["product"]
        pool = pools.get(product, [])
        if not pool:
            emp["assigned"] = []
            emp["matched"] = []
            emp["test_id"] = ""
            emp["portal_status"] = f"{emp['portal_status']} — no QB for {emp.get('raw_product') or product}"
            print(f"LOGIN ONLY {emp['employee_id']} {emp['full_name']} ({product} has no questions in QB-new)")
        else:
            assigned, remark = assign_display_questions(product, pool, employee_id=emp["employee_id"])
            assigned = [q for q in assigned if q]
            if len(assigned) < QUESTIONS_PER_EMPLOYEE:
                print(f"ERROR: {emp['employee_id']} only got {len(assigned)} {product} questions. {remark}")
                return 1
            matched = []
            for q_text in assigned:
                item = bank.get(normalize_question_key(q_text)) or bank.get(
                    normalize_question_key(q_text.split("] ", 1)[-1])
                )
                if not item:
                    print(f"ERROR: unmatched question for {emp['employee_id']}: {q_text[:120]}")
                    return 1
                matched.append({**item, "question_text": q_text})
            emp["assigned"] = assigned
            emp["matched"] = matched

        if email_taken_by_other(client, emp["email"], emp["employee_id"]):
            fallback = f"{emp['employee_id']}@iptel.portal.local"
            print(f"  Nokia email {emp['email']} already on another Emp ID; using {fallback}")
            emp["email"] = fallback

        existing_acc = account_map.get(emp["employee_id"].upper(), {})
        account = {
            **existing_acc,
            "employee_id": emp["employee_id"],
            "full_name": emp["full_name"],
            "email": emp["email"],
            "department": emp.get("domain") or "IPTEL",
            "role": "employee",
            "product": product,
            "assessment_only": True,
            "product_qb_eligible": bool(emp.get("matched")),
            "is_first_login": True,
            "password_hash": emp["password_hash"],
            "password_salt": emp["password_salt"],
            "xp_points": existing_acc.get("xp_points", 0) if emp["action"] != "remap" else 0,
            "streak_days": existing_acc.get("streak_days", 0) if emp["action"] != "remap" else 0,
            "skill_level": existing_acc.get("skill_level", "beginner"),
            "ai_readiness_score": existing_acc.get("ai_readiness_score", 0) if emp["action"] != "remap" else 0,
        }
        account_map[emp["employee_id"].upper()] = account
        emp["account"] = account

        profile_map[emp["employee_id"].upper()] = {
            "employee_id": emp["employee_id"],
            "full_name": emp["full_name"],
            "role": emp.get("role") or "employee",
            "domain": emp.get("domain") or "IPTEL",
            "product": product,
            "email": emp["email"],
            "ddh": emp.get("ddh") or "",
            "emp_status": "Confirmed",
            "remarks": REMARKS,
            "assigned_questions": emp.get("assigned") or [],
            "assigned_question_count": len(emp.get("assigned") or []),
        }
        if emp.get("assigned"):
            update_mapping_row(emp, emp["assigned"])

        if emp.get("matched"):
            test_id = str(uuid.uuid4())
            emp["test_id"] = test_id
            new_tests.append(
                {
                    "id": test_id,
                    "employee_id": emp["employee_id"],
                    "employee_code": emp["employee_id"],
                    "topic_id": TOPIC_ID,
                    "subject_id": SUBJECT_ID,
                    "difficulty": "medium",
                    "total_questions": len(emp["matched"]),
                    "time_limit_seconds": TIME_LIMIT_SECONDS,
                    "status": "pending",
                    "current_question_index": 0,
                    "started_at": None,
                    "completed_at": None,
                    "in_progress": None,
                    "created_at": now,
                    "topic_title": product,
                    "subject_title": SUBJECT_TITLE,
                }
            )
            for idx, q in enumerate(emp["matched"]):
                new_questions.append(
                    {
                        "id": str(uuid.uuid4()),
                        "test_id": test_id,
                        "question_index": idx,
                        "question_text": q["question_text"],
                        "options": q["options"],
                        "correct_option_index": q["correct_option_index"],
                        "explanation": q["explanation"],
                        "difficulty": q["difficulty"],
                        "topic_id": TOPIC_ID,
                        "topic_title": q.get("category") or q["topic_title"],
                        "created_at": now,
                    }
                )
            new_manifest.append(
                {
                    "employee_id": emp["employee_id"],
                    "full_name": emp["full_name"],
                    "product": product,
                    "test_id": test_id,
                    "question_count": len(emp["matched"]),
                    "missing_questions": [],
                }
            )
        print(
            f"Prepared {emp['employee_id']} {emp['previous_name']} -> {emp['full_name']} "
            f"{product} ({emp['action']}, {len(emp.get('matched') or [])} Qs)"
        )

    accounts["employees"] = list(account_map.values())
    save_json(ACCOUNTS_FILE, accounts)
    save_json(PROFILES_FILE, list(profile_map.values()))
    append_credentials_rows(to_add)
    write_share_file(share_rows)
    append_local_tests(new_tests, new_questions, new_manifest)

    print(f"\nSyncing {len(to_add)} portal users to Supabase...")
    for emp in to_add:
        employee_uuid = upsert_supabase_employee(client, emp, emp["account"])
        if emp.get("matched") and emp.get("test_id"):
            if emp["action"] == "remap":
                replace_or_create_test(client, employee_uuid, emp, emp["matched"], emp["test_id"])
            else:
                create_test_if_missing(client, employee_uuid, emp, emp["matched"], emp["test_id"])
        print(f"  synced {emp['employee_id']} {emp['full_name']}")

    print("\n=== Employee Portal logins ===")
    print("URL: https://ai-interview-ics-poc.vercel.app/employee")
    for emp in share_rows:
        print(
            f"{emp['employee_id']}  {emp['full_name']}  {emp['password']}  "
            f"{emp.get('raw_product') or emp['product']}  {emp['portal_status']}"
        )
    print(f"\nShare file: {SHARE_FILE_ROOT}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
