"""
Create Employee Portal credentials for the 16 Sep 2026 PACO list.

Does not wipe other portal users. Same Emp ID + different name is remapped
(ID reuse). Same person already on a Q-form is left unchanged.
"""
from __future__ import annotations

import re
import sys
import uuid
from datetime import datetime, timezone
from pathlib import Path

import requests
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))

from add_paco_cmg3_portal_users import (  # noqa: E402
    email_taken_by_other,
    replace_or_create_test,
)
from add_paco_phase2_portal_users import (  # noqa: E402
    ACCOUNTS_FILE,
    PROFILES_FILE,
    QUESTIONS_PER_EMPLOYEE,
    SUBJECT_ID,
    SUBJECT_TITLE,
    TIME_LIMIT_SECONDS,
    append_credentials_rows,
    append_local_tests,
    assign_display_questions,
    build_bank_index,
    clean,
    generate_password,
    hash_password,
    load_existing_passwords,
    load_json,
    normalize_question_key,
    parse_qb_new_xlsx,
    save_json,
    update_mapping_row,
    upsert_supabase_employee,
)
import add_paco_phase2_portal_users as phase2  # noqa: E402
from reassign_test_questions_supabase import TOPIC_ID, SupabaseClient, load_env  # noqa: E402

DOCS_CANDIDATES = [
    ROOT / "NON-Needed docs",
    ROOT / "docs" / "NON-Needed docs",
]
DOCS = next((p for p in DOCS_CANDIDATES if (p / "QB-new.xlsx").exists()), DOCS_CANDIDATES[0])
QB_FILE = DOCS / "QB-new.xlsx"
SHARE_FILE = DOCS / "PACO_16Sep26_User_Credentials.xlsx"
phase2.QB_FILE = QB_FILE
phase2.CREDENTIALS_FILE = DOCS / "Employee_User_Credentials.xlsx"
phase2.MAPPING_FILE = DOCS / "Resource_Question_Mapping.xlsx"
phase2.DOCS = DOCS
REMARKS = "PACO credentials 16 Sep 2026"

PEOPLE = [
    {"employee_id": "1042767", "full_name": "Nalluri Prabhuteja", "email": "nalluri.prabhu_teja.ext@nokia.com", "product": "CMG"},
    {"employee_id": "1039687", "full_name": "Chevuri Venkata Dhathrija", "email": "chevuri.dhathrija.ext@nokia.com", "product": "CMG"},
    {"employee_id": "1035294", "full_name": "Swethasri R", "email": "swethasri.ravikumar.ext@nokia.com", "product": "CMG"},
    {"employee_id": "1035545", "full_name": "Dhivya Sri Kathirvel", "email": "dhivya.kathirvel.ext@nokia.com", "product": "CMG"},
    {"employee_id": "1042385", "full_name": "Pakki Sai Sireesha", "email": "sai_sireesha.pakki.ext@nokia.com", "product": "CMG"},
    {"employee_id": "1041269", "full_name": "Shendkar Jayashri Babaji", "email": "jayashri.shendkar.ext@nokia.com", "product": "CMG"},
    {"employee_id": "1041631", "full_name": "Apoorva Sharma", "email": "apoorva.sharma.ext@nokia.com", "product": "CMG"},
    {"employee_id": "1042145", "full_name": "Vishwa Keerthi G", "email": "keerthi.ganamur.ext@nokia.com", "product": "CMG"},
    {"employee_id": "1043028", "full_name": "Amit Vyas", "email": "amit.vyas.ext@nokia.com", "product": "CMG"},
    {"employee_id": "1043025", "full_name": "Mohammed Sahmad Hussain", "email": "mohammed.sahmad_hussain.ext@nokia.com", "product": "CMG"},
    {"employee_id": "1043077", "full_name": "Ishan Gupta", "email": "ishan.1.gupta.ext@nokia.com", "product": "CMG"},
    {"employee_id": "1043156", "full_name": "Ritesh Srivastava", "email": "ritesh.srivastava.ext@nokia.com", "product": "CMG"},
    {"employee_id": "1043155", "full_name": "Mark G", "email": "mark.sigamani.ext@nokia.com", "product": "CMG"},
    {"employee_id": "1043138", "full_name": "Nikhil Kumar", "email": "nikhil.1.kumar.ext@nokia.com", "product": "CMG"},
    {"employee_id": "1043250", "full_name": "Shashi Ranjan Singh", "email": "shashi.r1.singh.ext@nokia.com", "product": "CMG"},
    {"employee_id": "1043320", "full_name": "Abhay Thakur", "email": "abhay.thakur.ext@nokia.com", "product": "CMG"},
    {"employee_id": "1043318", "full_name": "Vinod Kumar", "email": "vinod.2.kumar.ext@nokia.com", "product": "CMG"},
    {"employee_id": "1043437", "full_name": "S Sivachandiran", "email": "sivachandiran.s.ext@nokia.com", "product": "CMG"},
    {"employee_id": "1043494", "full_name": "Himanshi Redhu", "email": "himanshi.redhu.ext@nokia.com", "product": "CMG"},
    {"employee_id": "1029007", "full_name": "Subathra T", "email": "subathra.t.ext@nokia.com", "product": "CMM"},
    {"employee_id": "1040284", "full_name": "Nikitha Bhakre", "email": "nikitha.bhakre.ext@nokia.com", "product": "CMM"},
    {"employee_id": "1042355", "full_name": "Vishal Pratap Singh", "email": "vishal.pratap.ext@nokia.com", "product": "CMM"},
    {"employee_id": "1042661", "full_name": "Sunil Kumar", "email": "sunil.8.kumar.ext@nokia.com", "product": "CMM"},
    {"employee_id": "1042957", "full_name": "Omprakash Kumar", "email": "omprakash.kumar.ext@nokia.com", "product": "CMM"},
    {"employee_id": "1043165", "full_name": "Sriram Azhakia Nayagam", "email": "sriram.azhakia_nayagam.ext@nokia.com", "product": "CMM"},
    {"employee_id": "1043453", "full_name": "Vinod 10 Kumar", "email": "Vinod_10.kumar.ext@nokia.com", "product": "CMM"},
    {"employee_id": "1043060", "full_name": "Rohit Pratap Singh", "email": "rohit.p.singh.ext@nokia.com", "product": "NRD"},
    {"employee_id": "1042997", "full_name": "Ravi Ranjan Jha", "email": "raviranjan.kumar.ext@nokia.com", "product": "CMM"},
    {"employee_id": "1042153", "full_name": "Avinash Kumar Sharma", "email": "avinash.k.sharma.ext@nokia.com", "product": "CMM"},
]


def name_tokens(value: str) -> set[str]:
    return {t for t in re.findall(r"[a-z]{2,}", (value or "").lower())}


def names_match(left: str, right: str) -> bool:
    a = name_tokens(left)
    b = name_tokens(right)
    if not a or not b:
        return False
    return bool(a & b)


def write_share_file(employees: list[dict]) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "PACO 16 Sep 26"
    header = [
        "Emp ID",
        "Employee Name",
        "Initial Password",
        "Nokia Email ID",
        "Role",
        "Domain",
        "Product",
        "Login ID",
        "Portal status",
    ]
    ws.append(header)
    for emp in employees:
        ws.append(
            [
                emp["employee_id"],
                emp["full_name"],
                emp["password"],
                emp["email"],
                "employee",
                "PACO",
                emp["product"],
                emp["employee_id"],
                emp["portal_status"],
            ]
        )
        row = ws.max_row
        ws.cell(row, 1).value = str(emp["employee_id"])
        ws.cell(row, 8).value = str(emp["employee_id"])
        ws.cell(row, 3).font = Font(name="Consolas", size=10, bold=True, color="4338CA")

    header_fill = PatternFill("solid", fgColor="4F46E5")
    for cell in ws[1]:
        cell.font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(vertical="center", horizontal="center")
    ws.row_dimensions[1].height = 28
    widths = [14, 32, 22, 42, 12, 10, 10, 14, 28]
    for idx, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = width

    note = wb.create_sheet("How to login")
    note["A1"] = "Employee Portal login"
    note["A1"].font = Font(bold=True, size=14)
    note["A3"] = "URL"
    note["B3"] = "https://ai-interview-ics-poc.vercel.app/employee"
    note["A4"] = "Login ID"
    note["B4"] = "Emp ID from this sheet"
    note["A5"] = "Password"
    note["B5"] = "Initial Password column"
    note["A7"] = "Newly added / remapped = new first-login password. Already in portal = existing password, tests were not changed."
    note["A8"] = "Treat this file as confidential."
    note.column_dimensions["A"].width = 16
    note.column_dimensions["B"].width = 100
    SHARE_FILE.parent.mkdir(parents=True, exist_ok=True)
    wb.save(SHARE_FILE)
    print(f"Wrote share file: {SHARE_FILE} ({len(employees)} users)")


def raise_for_status(resp, action: str) -> None:
    if resp.ok:
        return
    raise RuntimeError(f"{action} failed {resp.status_code}: {resp.text}")


def create_test_if_missing(client, employee_uuid: str, emp: dict, matched: list[dict], test_id: str) -> None:
    existing = client.get(
        "tests",
        {
            "employee_id": f"eq.{employee_uuid}",
            "topic_id": f"eq.{TOPIC_ID}",
            "select": "id,status",
            "limit": "1",
        },
    )
    if existing:
        print(f"  skip existing test {existing[0]['id']} status={existing[0].get('status')}")
        return
    replace_or_create_test(client, employee_uuid, emp, matched, test_id)


def main() -> int:
    if not QB_FILE.exists():
        print(f"ERROR: missing {QB_FILE}")
        return 1

    stored_passwords = load_existing_passwords()
    accounts = load_json(ACCOUNTS_FILE, {"employees": []})
    account_map = {
        clean(acc.get("employee_id")).upper(): acc
        for acc in accounts.get("employees", [])
        if acc.get("employee_id")
    }
    profiles = load_json(PROFILES_FILE, [])
    profile_map = {clean(row.get("employee_id")).upper(): row for row in profiles}

    url, key = load_env()
    client = SupabaseClient(url, key)

    to_add = []
    share_rows = []
    for raw in PEOPLE:
        emp = {
            **raw,
            "role": "employee",
            "domain": "PACO",
            "ddh": "",
        }
        existing_acc = account_map.get(emp["employee_id"].upper(), {})
        existing_name = clean(existing_acc.get("full_name"))
        eligible = bool(existing_acc.get("product_qb_eligible"))
        same_person = (not existing_name) or names_match(existing_name, emp["full_name"])
        if existing_name and not same_person:
            emp["action"] = "remap"
            emp["previous_name"] = existing_name
        elif eligible and same_person:
            emp["action"] = "skip"
            emp["previous_name"] = existing_name or emp["full_name"]
        else:
            emp["action"] = "add"
            emp["previous_name"] = existing_name or "(new)"

        if emp["action"] == "skip":
            emp["password"] = stored_passwords.get(emp["employee_id"].upper()) or "(already in portal — password unchanged)"
            emp["portal_status"] = "Already in portal"
            emp["skip"] = True
            share_rows.append(emp)
            print(f"SKIP {emp['employee_id']} {emp['full_name']} (already Q-form eligible)")
            continue

        emp["password"] = stored_passwords.get(emp["employee_id"].upper()) or generate_password(
            emp["full_name"], emp["employee_id"]
        )
        emp["password_hash"], emp["password_salt"] = hash_password(emp["password"])
        emp["portal_status"] = "Remapped Emp ID" if emp["action"] == "remap" else "Newly added"
        emp["skip"] = False
        to_add.append(emp)
        share_rows.append(emp)

    print(f"Already in portal: {sum(1 for e in share_rows if e['skip'])}")
    print(f"To add/remap: {len(to_add)}")

    pools, mcq_records = parse_qb_new_xlsx(QB_FILE)
    bank = build_bank_index(mcq_records)
    now = datetime.now(timezone.utc).isoformat()
    new_tests = []
    new_questions = []
    new_manifest = []

    for emp in to_add:
        product = emp["product"]
        pool = pools.get(product, [])
        assigned, remark = assign_display_questions(product, pool, employee_id=emp["employee_id"])
        assigned = [q for q in assigned if q]
        if len(assigned) < QUESTIONS_PER_EMPLOYEE:
            print(f"ERROR: {emp['employee_id']} only got {len(assigned)} {product} questions. {remark}")
            return 1

        matched = []
        for q_text in assigned:
            item = bank.get(normalize_question_key(q_text)) or bank.get(
                normalize_question_key(q_text.split("] ", 1)[-1])
            )
            if not item:
                print(f"ERROR: unmatched question for {emp['employee_id']}: {q_text[:120]}")
                return 1
            matched.append({**item, "question_text": q_text})

        if email_taken_by_other(client, emp["email"], emp["employee_id"]):
            fallback = f"{emp['employee_id']}@paco.portal.local"
            print(f"  Nokia email {emp['email']} already on another Emp ID; using {fallback}")
            emp["email"] = fallback

        existing_acc = account_map.get(emp["employee_id"].upper(), {})
        account = {
            **existing_acc,
            "employee_id": emp["employee_id"],
            "full_name": emp["full_name"],
            "email": emp["email"],
            "department": "PACO",
            "role": "employee",
            "product": product,
            "assessment_only": True,
            "product_qb_eligible": True,
            "is_first_login": True,
            "password_hash": emp["password_hash"],
            "password_salt": emp["password_salt"],
            "xp_points": existing_acc.get("xp_points", 0) if emp["action"] != "remap" else 0,
            "streak_days": existing_acc.get("streak_days", 0) if emp["action"] != "remap" else 0,
            "skill_level": existing_acc.get("skill_level", "beginner"),
            "ai_readiness_score": existing_acc.get("ai_readiness_score", 0) if emp["action"] != "remap" else 0,
        }
        account_map[emp["employee_id"].upper()] = account
        emp["account"] = account
        emp["assigned"] = assigned
        emp["matched"] = matched

        profile_map[emp["employee_id"].upper()] = {
            "employee_id": emp["employee_id"],
            "full_name": emp["full_name"],
            "role": "employee",
            "domain": "PACO",
            "product": product,
            "email": emp["email"],
            "ddh": "",
            "emp_status": "Confirmed",
            "remarks": REMARKS,
            "assigned_questions": assigned,
            "assigned_question_count": len(assigned),
        }
        update_mapping_row(emp, assigned)

        test_id = str(uuid.uuid4())
        emp["test_id"] = test_id
        new_tests.append(
            {
                "id": test_id,
                "employee_id": emp["employee_id"],
                "employee_code": emp["employee_id"],
                "topic_id": TOPIC_ID,
                "subject_id": SUBJECT_ID,
                "difficulty": "medium",
                "total_questions": len(matched),
                "time_limit_seconds": TIME_LIMIT_SECONDS,
                "status": "pending",
                "current_question_index": 0,
                "started_at": None,
                "completed_at": None,
                "in_progress": None,
                "created_at": now,
                "topic_title": product,
                "subject_title": SUBJECT_TITLE,
            }
        )
        for idx, q in enumerate(matched):
            new_questions.append(
                {
                    "id": str(uuid.uuid4()),
                    "test_id": test_id,
                    "question_index": idx,
                    "question_text": q["question_text"],
                    "options": q["options"],
                    "correct_option_index": q["correct_option_index"],
                    "explanation": q["explanation"],
                    "difficulty": q["difficulty"],
                    "topic_id": TOPIC_ID,
                    "topic_title": q.get("category") or q["topic_title"],
                    "created_at": now,
                }
            )
        new_manifest.append(
            {
                "employee_id": emp["employee_id"],
                "full_name": emp["full_name"],
                "product": product,
                "test_id": test_id,
                "question_count": len(matched),
                "missing_questions": [],
            }
        )
        print(
            f"Prepared {emp['employee_id']} {emp['previous_name']} -> {emp['full_name']} "
            f"{product} ({emp['action']}, {len(matched)} Qs)"
        )

    accounts["employees"] = list(account_map.values())
    save_json(ACCOUNTS_FILE, accounts)
    save_json(PROFILES_FILE, list(profile_map.values()))
    append_credentials_rows(to_add)
    write_share_file(share_rows)
    append_local_tests(new_tests, new_questions, new_manifest)

    print(f"\nSyncing {len(to_add)} portal users to Supabase...")
    for emp in to_add:
        employee_uuid = upsert_supabase_employee(client, emp, emp["account"])
        if emp["action"] == "remap":
            replace_or_create_test(client, employee_uuid, emp, emp["matched"], emp["test_id"])
        else:
            create_test_if_missing(client, employee_uuid, emp, emp["matched"], emp["test_id"])
        print(f"  synced {emp['employee_id']} {emp['full_name']}")

    print("\n=== Employee Portal logins ===")
    print("URL: https://ai-interview-ics-poc.vercel.app/employee")
    for emp in share_rows:
        print(f"{emp['employee_id']}  {emp['full_name']}  {emp['password']}  {emp['product']}  {emp['portal_status']}")
    print(f"\nShare file: {SHARE_FILE}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
